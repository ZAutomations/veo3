#!/usr/bin/env python
"""
Tests for story_to_excel.py.

The workbook itself is asserted by reading it back with openpyxl, so a change
that silently stops writing a sheet fails here rather than in the user's hands.
The most important assertion is the last one: an export must never modify the
story JSON, because the engine builds from that file immediately afterwards.

Run: python test_story_to_excel.py
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import story_to_excel as S


def story_fixture():
    return {
        "title": "The Price of Obligation",
        "description": "A marriage under a debt.",
        "niche": "relationship",
        "moral": "a debt is not a marriage",
        "target_audience": "adults",
        "style": "dim, warm, intimate",
        "place": "a bedroom",
        "blocking": "she reads, he stands",
        "aspect_ratio": "16:9",
        "scene_seconds": 8,
        "total_scenes": 2,
        "video_duration": "16 seconds",
        "narrated": False,
        "narration_scope": "none",
        "silent_cast": ["godwin"],
        "scenes": [
            {
                "_scene_number": 1,
                "_scene_title": "The Nightly Request",
                "_timing": "0:00-0:08",
                "scene_builder_action": "text_to_video",
                "extend_from_last_frame": False,
                "characters": ["Tari", "Godwin"],
                "script_line": "",
                "dialogue": [
                    {"speaker": "Tari", "line": "Did you marry me to have a debtor?"},
                    {"speaker": "Godwin", "line": "Come to bed, Tari."},
                ],
                "narrative_context": "A dim bedroom.",
                "veo3_prompt": "[SHOT] A cozy, dimly lit bedroom.",
            },
            {
                "_scene_number": 2,
                "_scene_title": "The Ledger",
                "_timing": "0:08-0:16",
                "scene_builder_action": "extend",
                "extend_from_last_frame": True,
                "characters": ["Godwin"],
                "dialogue": "A single line, not a list.",
                "veo3_prompt": "[SHOT] He opens the ledger.",
            },
        ],
        "character_descriptions": {"godwin": "A 34-year-old man.", "tari": "A woman of 30."},
        "character_references": {"godwin": "character_refs/godwin.jpg",
                                 "tari": "character_refs/tari.jpg"},
    }


class ExcelExport(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="veo3-xlsx-")
        self.story = os.path.join(self.dir, "the_price_of_obligation_story.json")
        with open(self.story, "w", encoding="utf-8") as fh:
            json.dump(story_fixture(), fh, indent=2)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    # ── loading ─────────────────────────────────────────────────────────────
    def test_loads_a_story_file(self):
        path, data = S.load_story(self.story)
        self.assertEqual(path, self.story)
        self.assertEqual(data["title"], "The Price of Obligation")

    def test_loads_from_the_folder(self):
        path, data = S.load_story(self.dir)
        self.assertTrue(path.endswith("_story.json"))
        self.assertEqual(data["total_scenes"], 2)

    def test_missing_story_exits(self):
        with self.assertRaises(SystemExit):
            S.load_story(os.path.join(self.dir, "nope.json"))

    def test_bad_json_exits(self):
        bad = os.path.join(self.dir, "broken_story.json")
        with open(bad, "w", encoding="utf-8") as fh:
            fh.write("{not json")
        with self.assertRaises(SystemExit):
            S.load_story(bad)

    def test_empty_folder_exits(self):
        empty = tempfile.mkdtemp(prefix="veo3-empty-")
        try:
            with self.assertRaises(SystemExit):
                S.load_story(empty)
        finally:
            shutil.rmtree(empty, ignore_errors=True)

    # ── shaping ─────────────────────────────────────────────────────────────
    def test_as_text_covers_the_shapes_a_story_uses(self):
        self.assertEqual(S.as_text(None), "")
        self.assertEqual(S.as_text("x"), "x")
        self.assertEqual(S.as_text(True), "yes")
        self.assertEqual(S.as_text(False), "no")
        self.assertEqual(S.as_text(8), 8)
        self.assertEqual(S.as_text(["a", "b"]), "a\nb")
        self.assertEqual(S.as_text({"k": "v"}), "k: v")

    def test_dialogue_is_read_from_all_three_shapes(self):
        self.assertEqual(S.dialogue_text([{"speaker": "A", "line": "one"}]), "A: one")
        self.assertEqual(S.dialogue_text("plain"), "plain")
        self.assertEqual(S.dialogue_text({"speaker": "A", "line": "one"}), "A: one")
        self.assertEqual(S.dialogue_text({"A": "one", "B": "two"}), "A: one\nB: two")
        self.assertEqual(S.dialogue_text([]), "")
        self.assertEqual(S.dialogue_text(None), "")
        # A line with no speaker still appears rather than being dropped.
        self.assertEqual(S.dialogue_text([{"line": "bare"}]), "bare")

    def test_dialogue_rows_are_one_per_line(self):
        rows = S.dialogue_rows(story_fixture())
        self.assertEqual(len(rows), 3)          # 2 in scene 1, 1 in scene 2
        self.assertEqual(rows[0], (1, "The Nightly Request", "Tari",
                                   "Did you marry me to have a debtor?"))
        self.assertEqual(rows[2][0], 2)         # scene number, not list index
        self.assertEqual(rows[2][3], "A single line, not a list.")

    def test_cast_rows_pair_name_with_sheet(self):
        rows = dict((r[0], r) for r in S.cast_rows(story_fixture()))
        self.assertIn("godwin", rows)
        self.assertEqual(rows["godwin"][1], "A 34-year-old man.")
        self.assertEqual(rows["godwin"][2], "godwin.jpg")

    def test_cast_tolerates_a_story_without_a_cast(self):
        self.assertEqual(S.cast_rows({}), [])

    # ── the workbook ────────────────────────────────────────────────────────
    def test_writes_four_sheets(self):
        out = os.path.join(self.dir, "book.xlsx")
        S.build_workbook(story_fixture(), out)
        self.assertTrue(os.path.exists(out))
        from openpyxl import load_workbook
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Story", "Scenes", "Dialogue", "Cast"])

    def test_scenes_sheet_has_a_row_per_scene_and_a_frozen_header(self):
        out = os.path.join(self.dir, "book.xlsx")
        S.build_workbook(story_fixture(), out)
        from openpyxl import load_workbook
        ws = load_workbook(out)["Scenes"]
        self.assertEqual(ws.max_row, 3)                 # header + 2 scenes
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual([c.value for c in ws[1]][0], "Scene")
        self.assertEqual(ws.cell(row=2, column=1).value, 1)
        self.assertEqual(ws.cell(row=3, column=2).value, "The Ledger")

    def test_the_dialogue_column_is_readable_not_a_repr(self):
        out = os.path.join(self.dir, "book.xlsx")
        S.build_workbook(story_fixture(), out)
        from openpyxl import load_workbook
        ws = load_workbook(out)["Scenes"]
        cell = ws.cell(row=2, column=8).value      # "Dialogue"
        self.assertIn("Tari: Did you marry me", cell)
        self.assertNotIn("{", cell)
        self.assertNotIn("[", cell)

    def test_story_sheet_lists_the_metadata(self):
        out = os.path.join(self.dir, "book.xlsx")
        S.build_workbook(story_fixture(), out)
        from openpyxl import load_workbook
        ws = load_workbook(out)["Story"]
        labels = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        self.assertIn("Title", labels)
        self.assertIn("Niche", labels)
        self.assertIn("Silent cast", labels)

    def test_a_story_with_no_scenes_still_writes_a_workbook(self):
        out = os.path.join(self.dir, "empty.xlsx")
        S.build_workbook({"title": "Bare"}, out)
        from openpyxl import load_workbook
        wb = load_workbook(out)
        self.assertEqual(wb["Scenes"].max_row, 1)   # the header alone

    # ── the CLI ─────────────────────────────────────────────────────────────
    def test_cli_writes_beside_the_story_by_default(self):
        r = subprocess.run([sys.executable, os.path.join(HERE, "story_to_excel.py"), self.story],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        expected = os.path.join(self.dir, "the_price_of_obligation.xlsx")
        self.assertTrue(os.path.exists(expected), r.stdout)
        self.assertIn("2 scene(s)", r.stdout)

    def test_cli_honours_an_output_path(self):
        out = os.path.join(self.dir, "named.xlsx")
        r = subprocess.run([sys.executable, os.path.join(HERE, "story_to_excel.py"),
                            self.story, "-o", out], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertTrue(os.path.exists(out))

    def test_cli_reports_a_missing_story_without_a_traceback(self):
        r = subprocess.run([sys.executable, os.path.join(HERE, "story_to_excel.py"),
                            os.path.join(self.dir, "nope.json")], capture_output=True, text=True)
        self.assertEqual(r.returncode, 1)
        self.assertIn("No such story", r.stderr)
        self.assertNotIn("Traceback", r.stderr)

    # ── the promise ─────────────────────────────────────────────────────────
    def test_the_export_never_modifies_the_story(self):
        with open(self.story, "rb") as fh:
            before = fh.read()
        r = subprocess.run([sys.executable, os.path.join(HERE, "story_to_excel.py"), self.story],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        with open(self.story, "rb") as fh:
            after = fh.read()
        self.assertEqual(before, after, "the exporter rewrote the story JSON")


if __name__ == "__main__":
    unittest.main(verbosity=2)
