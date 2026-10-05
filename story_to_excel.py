#!/usr/bin/env python
"""
STORY -> EXCEL
==============
One .xlsx per story, for reading, reviewing and editing outside the pipeline.
A story JSON is fine for a machine and painful for a person: fourteen scenes of
prompt text and dialogue in one nested object, with the scene you want to fix
somewhere in the middle of it. This lays the same data out as sheets.

  Story     the metadata - title, niche, moral, style, place, ratio, timing
  Scenes    one row per scene; the sheet you actually work from
  Dialogue  one row per line with its speaker, for a voice or subtitle pass
  Cast      each character, its description, and the sheet it is built from

Usage:
    python story_to_excel.py <story.json> [-o out.xlsx]
    python story_to_excel.py <story-folder>

With no -o the workbook is written beside the story as
<story name>.xlsx. The story JSON is only ever READ - nothing here writes to it,
so an export can never damage the story the engine is about to build.
"""

import argparse
import json
import os
import re
import sys

# Column layouts: (heading, key, width, wrap). Order here is the sheet order.
SCENE_COLUMNS = [
    ("Scene", "_scene_number", 7, False),
    ("Title", "_scene_title", 26, True),
    ("Timing", "_timing", 13, False),
    ("Action", "scene_builder_action", 17, False),
    ("Extend", "extend_from_last_frame", 8, False),
    ("Characters", "characters", 20, True),
    ("Script line", "script_line", 30, True),
    ("Dialogue", "dialogue", 46, True),
    ("Narrative context", "narrative_context", 50, True),
    ("Veo3 prompt", "veo3_prompt", 62, True),
]

STORY_FIELDS = [
    ("Title", "title"),
    ("Description", "description"),
    ("Niche", "niche"),
    ("Moral", "moral"),
    ("Target audience", "target_audience"),
    ("Style", "style"),
    ("Place", "place"),
    ("Blocking", "blocking"),
    ("Aspect ratio", "aspect_ratio"),
    ("Seconds per scene", "scene_seconds"),
    ("Total scenes", "total_scenes"),
    ("Duration", "video_duration"),
    ("Narrated", "narrated"),
    ("Narration scope", "narration_scope"),
    ("Silent cast", "silent_cast"),
]


def die(msg):
    print(msg, file=sys.stderr)
    sys.exit(1)


def load_story(path):
    """The story JSON at `path`, or in the folder `path` if it is a folder."""
    if os.path.isdir(path):
        cands = [f for f in os.listdir(path) if f.endswith("_story.json")] or \
                [f for f in os.listdir(path) if f.endswith(".json")]
        if not cands:
            die(f"No story JSON in {path}")
        # A folder can hold refs.json and content maps too; a *_story.json wins.
        story = [c for c in cands if c.endswith("_story.json")] or cands
        path = os.path.join(path, sorted(story)[0])
    if not os.path.isfile(path):
        die(f"No such story: {path}")
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except json.JSONDecodeError as e:
        die(f"{path} is not valid JSON: {e}")
    if not isinstance(data, dict):
        die(f"{path} does not hold a story object")
    return path, data


def as_text(v):
    """A cell value. Lists become lines, dicts become 'k: v' lines."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, list):
        return "\n".join(as_text(x) for x in v)
    if isinstance(v, dict):
        return "\n".join(f"{k}: {as_text(x)}" for k, x in v.items())
    return str(v)


def dialogue_text(d):
    """The scene's dialogue as readable lines, whatever shape it arrived in."""
    if not d:
        return ""
    if isinstance(d, str):
        return d
    if isinstance(d, dict):
        # A dict of speaker -> line, or a single {speaker, line}.
        if "line" in d:
            return f"{d.get('speaker', '')}: {d['line']}".strip(": ")
        return "\n".join(f"{k}: {as_text(v)}" for k, v in d.items())
    if isinstance(d, list):
        out = []
        for item in d:
            if isinstance(item, dict):
                spk = str(item.get("speaker", "") or "").strip()
                line = as_text(item.get("line", item.get("text", ""))).strip()
                if spk and line:
                    out.append(f"{spk}: {line}")
                elif line:
                    out.append(line)
            else:
                out.append(as_text(item))
        return "\n".join(x for x in out if x)
    return str(d)


def dialogue_rows(story):
    """(scene no, title, speaker, line) for every spoken line in the story."""
    rows = []
    for i, sc in enumerate(story.get("scenes") or [], start=1):
        if not isinstance(sc, dict):
            continue
        num = sc.get("_scene_number", i)
        title = sc.get("_scene_title", "")
        d = sc.get("dialogue")
        if isinstance(d, str) and d.strip():
            rows.append((num, title, "", d.strip()))
        elif isinstance(d, dict):
            if "line" in d:
                rows.append((num, title, d.get("speaker", ""), as_text(d["line"])))
            else:
                for k, v in d.items():
                    rows.append((num, title, k, as_text(v)))
        elif isinstance(d, list):
            for item in d:
                if isinstance(item, dict):
                    rows.append((num, title,
                                 item.get("speaker", ""),
                                 as_text(item.get("line", item.get("text", "")))))
                else:
                    rows.append((num, title, "", as_text(item)))
    return rows


def cast_rows(story):
    """(name, description, sheet file, path) for every character."""
    descs = story.get("character_descriptions") or {}
    refs = story.get("character_references") or {}
    names = list(descs.keys()) + [n for n in refs.keys() if n not in descs]
    rows = []
    for n in names:
        p = refs.get(n) or ""
        rows.append((n, as_text(descs.get(n, "")),
                     os.path.basename(str(p)) if p else "", str(p)))
    return rows


def build_workbook(story, out_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        die("openpyxl is needed to write .xlsx. Install it with:\n"
            "    pip install openpyxl")

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")
    top_wrap = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top")

    wb = Workbook()

    def sheet(title):
        ws = wb.create_sheet(title)
        return ws

    def header(ws, headings, widths, wraps):
        for c, (h, w, wrap) in enumerate(zip(headings, widths, wraps), start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = top_wrap if wrap else top
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = "A2"

    # ── Story ───────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Story"
    ws.append(["Field", "Value"])
    for c in (1, 2):
        cell = ws.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = head_fill
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100
    for label, key in STORY_FIELDS:
        ws.append([label, as_text(story.get(key))])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
        row[1].alignment = top_wrap
        row[0].alignment = top
    ws.freeze_panes = "A2"

    # ── Scenes ──────────────────────────────────────────────────────────────
    ws = sheet("Scenes")
    header(ws, [c[0] for c in SCENE_COLUMNS], [c[2] for c in SCENE_COLUMNS],
           [c[3] for c in SCENE_COLUMNS])
    scenes = story.get("scenes") or []
    for i, sc in enumerate(scenes, start=1):
        if not isinstance(sc, dict):
            continue
        vals = []
        for _, key, _, _ in SCENE_COLUMNS:
            if key == "dialogue":
                vals.append(dialogue_text(sc.get("dialogue")))
            elif key == "_scene_number":
                vals.append(sc.get(key, i))
            elif key == "characters":
                vals.append(as_text(sc.get(key, "")))
            else:
                vals.append(as_text(sc.get(key, "")))
        ws.append(vals)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(SCENE_COLUMNS)):
        for c, (_, _, _, wrap) in zip(row, SCENE_COLUMNS):
            c.alignment = top_wrap if wrap else top

    # ── Dialogue ────────────────────────────────────────────────────────────
    ws = sheet("Dialogue")
    header(ws, ["Scene", "Title", "Speaker", "Line"], [7, 26, 18, 80],
           [False, True, False, True])
    for row in dialogue_rows(story):
        ws.append(list(row))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=4):
        row[1].alignment = top_wrap
        row[3].alignment = top_wrap

    # ── Cast ────────────────────────────────────────────────────────────────
    ws = sheet("Cast")
    header(ws, ["Character", "Description", "Reference sheet", "Path"],
           [20, 70, 34, 60], [False, True, True, True])
    for row in cast_rows(story):
        ws.append(list(row))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=4):
        for c in row:
            c.alignment = top_wrap

    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(
        description="Export a story JSON to a readable .xlsx workbook.")
    ap.add_argument("story", help="the story JSON, or the folder holding it")
    ap.add_argument("-o", "--out", help="output .xlsx (default: beside the story)")
    args = ap.parse_args()

    path, story = load_story(args.story)
    out = args.out
    if not out:
        stem = re.sub(r"_story$", "", os.path.splitext(os.path.basename(path))[0])
        out = os.path.join(os.path.dirname(os.path.abspath(path)), stem + ".xlsx")
    out = os.path.abspath(out)

    build_workbook(story, out)

    scenes = len(story.get("scenes") or [])
    lines = len(dialogue_rows(story))
    cast = len(cast_rows(story))
    print(f"Wrote {out}")
    print(f"  {scenes} scene(s), {lines} dialogue line(s), {cast} character(s)")
    print(f"  sheets: Story, Scenes, Dialogue, Cast")


if __name__ == "__main__":
    main()
